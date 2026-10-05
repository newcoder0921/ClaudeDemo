"""Run a product's pytest suite and write a unit test case document (Markdown + Excel) from the real results.

    python -m tools.build_unit_test_cases member_360
    python -m tools.build_unit_test_cases fraud_case
"""
from __future__ import annotations

import argparse
import ast
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[1]
FRAMEWORK = ("Framework", "-")


@dataclass(frozen=True)
class Suite:
    label: str
    cwd: Path
    test_path: str
    stories: dict[str, tuple[str, str]]   # test file stem -> (story, Jira key)


@dataclass(frozen=True)
class Product:
    code: str
    title: str
    out_dir: Path
    suites: tuple[Suite, ...] = field(default_factory=tuple)


TOOLS_SUITE = Suite("tools", PROJECT_ROOT, "tools/tests",
                    {"test_account_id_map": ("Shared: account ID mapping reference table", "SCRUM-32")})

PRODUCTS = {
    "member_360": Product("M360", "DP_Member_360", PROJECT_ROOT / "output" / "DP_Member_360", (
        Suite("product", PROJECT_ROOT / "output" / "DP_Member_360" / "04_code", "tests", {
            "test_ingest": ("FR-01 Ingest sources", "SCRUM-18"),
            "test_standardize": ("FR-02 Standardize data types", "SCRUM-19"),
            "test_types": ("FR-02 Standardize data types", "SCRUM-19"),
            "test_member_id": ("FR-03 Conform member ID", "SCRUM-20"),
            "test_build_core": ("FR-04 Build Member 360 core", "SCRUM-21"),
            "test_pending_attributes": ("FR-05 Integrate new feeds", "SCRUM-22"),
            "test_dq": ("FR-06 Data quality and exceptions", "SCRUM-23"),
            "test_publish": ("FR-07 Publish and govern", "SCRUM-24"),
        }),
        TOOLS_SUITE,
    )),
    "fraud_case": Product("FC", "DP_Fraud_Case", PROJECT_ROOT / "output" / "DP_Fraud_Case", (
        Suite("product", PROJECT_ROOT / "output" / "DP_Fraud_Case" / "04_code", "tests", {
            "test_ingest": ("FR-01 Ingest case extract", "SCRUM-26"),
            "test_standardize": ("FR-02 Standardize data types", "SCRUM-27"),
            "test_types": ("FR-02 Standardize data types", "SCRUM-27"),
            "test_lifecycle": ("FR-03 Lifecycle and loss checks", "SCRUM-28"),
            "test_member_link": ("FR-04 Link to Member 360", "SCRUM-29"),
            "test_metrics": ("FR-05 Derived case metrics", "SCRUM-30"),
            "test_publish": ("FR-06 Publish and govern", "SCRUM-31"),
            "test_routing": ("FR-06 Publish and govern (routing check)", "SCRUM-31"),
            "test_status_history": ("FR-06 Publish and govern (status history)", "SCRUM-31"),
            "test_account_link": ("FR-07 Account ID mapping", "SCRUM-32"),
        }),
        TOOLS_SUITE,
    )),
}

HEADERS = ["TC ID", "Story", "Jira", "Test file", "Test case", "Description / expected result", "Result", "Time (s)"]


def humanize(name: str) -> str:
    base, _, params = name.partition("[")
    text = base.removeprefix("test_").replace("_", " ")
    text = text[:1].upper() + text[1:]
    return f"{text} [{params.rstrip(']')}]" if params else text


def docstrings(test_file: Path) -> dict[str, str]:
    tree = ast.parse(test_file.read_text(encoding="utf-8"))
    return {n.name: ast.get_docstring(n) or "" for n in ast.walk(tree) if isinstance(n, ast.FunctionDef)}


def run_suite(suite: Suite) -> list[dict]:
    with tempfile.TemporaryDirectory() as tmp:
        report = Path(tmp) / "junit.xml"
        subprocess.run([sys.executable, "-m", "pytest", "-q", "-p", "no:cacheprovider", suite.test_path,
                        f"--junitxml={report}"], cwd=suite.cwd, capture_output=True, text=True)
        root = ET.parse(report).getroot()
    rows = []
    for case in root.iter("testcase"):
        module = case.get("classname", "").split(".")[-1]
        result = ("Fail" if case.find("failure") is not None or case.find("error") is not None
                  else "Skipped" if case.find("skipped") is not None else "Pass")
        test_file = suite.cwd / suite.test_path / f"{module}.py"
        doc = docstrings(test_file).get(case.get("name").split("[")[0], "")
        story, jira = suite.stories.get(module, FRAMEWORK)
        rows.append({"Story": story, "Jira": jira, "Test file": f"{suite.test_path}/{module}.py",
                     "Test case": case.get("name"), "Description / expected result": doc or humanize(case.get("name")),
                     "Result": result, "Time (s)": round(float(case.get("time", 0)), 3)})
    return rows


def write_markdown(product: Product, rows: list[dict], path: Path) -> None:
    passed = sum(r["Result"] == "Pass" for r in rows)
    lines = [f"# Unit Test Cases – {product.title}", "",
             f"- **Generated:** {date.today().isoformat()} by `python -m tools.build_unit_test_cases`, from a real pytest run",
             f"- **Result:** {passed} / {len(rows)} passed"
             + (f", {sum(r['Result'] == 'Fail' for r in rows)} failed" if passed != len(rows) else ""),
             "- **Type:** automated pytest. Story-level tests run the pipeline on the real sample sources; "
             "Framework tests cover shared building blocks in isolation.", "",
             "## Summary by story", "", "| Story | Jira | Test cases | Passed |", "|---|---|---|---|"]
    by_story: dict[tuple[str, str], list[dict]] = {}
    for r in rows:
        by_story.setdefault((r["Story"], r["Jira"]), []).append(r)
    for (story, jira), items in by_story.items():
        lines.append(f"| {story} | {jira} | {len(items)} | {sum(i['Result'] == 'Pass' for i in items)} |")
    lines += ["", "## Test cases", "", "| " + " | ".join(HEADERS) + " |", "|" + "---|" * len(HEADERS)]
    lines += ["| " + " | ".join(str(r[h]).replace("|", "\\|") for h in HEADERS) + " |" for r in rows]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_excel(rows: list[dict], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Unit_Test_Cases"
    ws.append(HEADERS)
    for c in ws[1]:
        c.fill, c.font = PatternFill("solid", fgColor="1F4E79"), Font(bold=True, color="FFFFFF")
    for r in rows:
        ws.append([r[h] for h in HEADERS])
    for i, width in enumerate([12, 34, 11, 32, 48, 70, 9, 9], 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        result = row[6]
        result.font = Font(bold=True, color="007A33" if result.value == "Pass" else "C00000")
    ws.freeze_panes, ws.auto_filter.ref = "A2", ws.dimensions
    wb.save(path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("product", choices=sorted(PRODUCTS))
    args = parser.parse_args(argv)
    product = PRODUCTS[args.product]

    rows = [r for suite in product.suites for r in run_suite(suite)]
    # Product stories in FR order, then shared tooling, then framework; pytest order is kept within a story.
    rows.sort(key=lambda r: (2 if r["Story"] == FRAMEWORK[0] else 0 if r["Story"].startswith("FR-") else 1,
                             r["Story"][:5]))
    for i, r in enumerate(rows, 1):
        r["TC ID"] = f"UT-{product.code}-{i:03d}"
    stem = product.title.removeprefix("DP_")
    write_markdown(product, rows, product.out_dir / f"Unit_Test_Cases_{stem}.md")
    write_excel(rows, product.out_dir / f"Unit_Test_Cases_{stem}.xlsx")
    failed = sum(r["Result"] == "Fail" for r in rows)
    print(f"{product.title}: {len(rows)} test cases, {len(rows) - failed} passed, {failed} failed -> "
          f"{product.out_dir / f'Unit_Test_Cases_{stem}.md'}")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
